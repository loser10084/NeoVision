from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"D:\project\AI\ImageIdentification")
DOWNLOADS = Path(r"C:\Users\hu196\Downloads")
TEMPLATE_PATH = DOWNLOADS / "项目开发文档模版.docx"
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_PATH = OUTPUT_DIR / "灵犀智影-项目开发文档_6_7_8.docx"

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(10.5)


DB_TABLES = [
    {
        "title": "（1）医生用户表（doctor_user）",
        "caption": "表 6.1 医生用户表",
        "rows": [
            ("1", "id", "bigint", "20", "√", "主键", "医生主键 ID"),
            ("2", "name", "varchar", "64", "√", "", "医生姓名"),
            ("3", "hospital", "varchar", "128", "√", "", "所属医院"),
            ("4", "dept", "varchar", "64", "√", "", "所属科室"),
            ("5", "mobile", "varchar", "20", "√", "唯一", "手机号，也是登录账号"),
            ("6", "password_hash", "varchar", "255", "√", "", "密码哈希值"),
            ("7", "status", "tinyint", "1", "√", "", "账号状态，1 为启用，0 为停用"),
            ("8", "created_at", "datetime", "19", "√", "", "创建时间"),
            ("9", "updated_at", "datetime", "19", "√", "", "更新时间"),
        ],
    },
    {
        "title": "（2）患者信息表（patient）",
        "caption": "表 6.2 患者信息表",
        "rows": [
            ("1", "id", "bigint", "20", "√", "主键", "患者主键 ID"),
            ("2", "patient_no", "varchar", "32", "√", "唯一", "系统生成的患者编号"),
            ("3", "name", "varchar", "64", "√", "", "患者姓名"),
            ("4", "sex", "enum", "8", "√", "", "性别，男/女/其他"),
            ("5", "age", "tinyint", "3", "√", "", "年龄"),
            ("6", "stage", "varchar", "32", "", "", "肿瘤分期，如 T3N2M0"),
            ("7", "diagnosis", "varchar", "128", "", "", "临床诊断"),
            ("8", "attending_id", "bigint", "20", "", "外键", "主治医生 ID，关联 doctor_user.id"),
            ("9", "created_at", "datetime", "19", "√", "", "创建时间"),
            ("10", "updated_at", "datetime", "19", "√", "", "更新时间"),
        ],
    },
    {
        "title": "（3）影像序列表（patient_study）",
        "caption": "表 6.3 影像序列表",
        "rows": [
            ("1", "id", "bigint", "20", "√", "主键", "影像序列主键 ID"),
            ("2", "patient_id", "bigint", "20", "√", "外键", "所属患者 ID"),
            ("3", "study_uid", "varchar", "128", "√", "", "DICOM Study UID 或内部影像号"),
            ("4", "modality", "varchar", "16", "√", "", "影像模态，如 CT、MR、PET-CT"),
            ("5", "description", "varchar", "255", "", "", "序列说明"),
            ("6", "status", "varchar", "32", "√", "", "序列状态，如处理中、已完成"),
            ("7", "acquired_at", "datetime", "19", "", "", "采集或上传时间"),
            ("8", "created_at", "datetime", "19", "√", "", "创建时间"),
            ("9", "updated_at", "datetime", "19", "√", "", "更新时间"),
        ],
    },
    {
        "title": "（4）分割结果表（contour_result）",
        "caption": "表 6.4 分割结果表",
        "rows": [
            ("1", "id", "bigint", "20", "√", "主键", "轮廓结果主键 ID"),
            ("2", "patient_id", "bigint", "20", "√", "外键", "所属患者 ID"),
            ("3", "study_id", "bigint", "20", "√", "外键", "所属影像序列 ID"),
            ("4", "type", "enum", "8", "√", "", "轮廓类型，GTV/CTV/OAR"),
            ("5", "version", "int", "11", "√", "", "轮廓版本号"),
            ("6", "status", "varchar", "32", "√", "", "审核状态，如待确认、已确认"),
            ("7", "storage_path", "varchar", "255", "√", "", "轮廓文件存储路径"),
            ("8", "confidence_map", "varchar", "255", "", "", "置信度热力图路径"),
            ("9", "created_by", "bigint", "20", "", "", "结果生成者，可为 AI 或医生"),
            ("10", "created_at", "datetime", "19", "√", "", "创建时间"),
            ("11", "updated_at", "datetime", "19", "√", "", "更新时间"),
        ],
    },
    {
        "title": "（5）上传文件表（file_upload）",
        "caption": "表 6.5 上传文件表",
        "rows": [
            ("1", "id", "bigint", "20", "√", "主键", "文件记录主键 ID"),
            ("2", "patient_id", "bigint", "20", "√", "外键", "所属患者 ID"),
            ("3", "study_id", "bigint", "20", "", "外键", "所属影像序列 ID"),
            ("4", "file_type", "varchar", "32", "√", "", "文件类型，如 dicom、nrrd、mesh"),
            ("5", "file_path", "varchar", "255", "√", "", "对象存储或文件服务器路径"),
            ("6", "size_bytes", "bigint", "20", "", "", "文件大小，单位字节"),
            ("7", "uploader_id", "bigint", "20", "", "", "上传者 ID"),
            ("8", "created_at", "datetime", "19", "√", "", "创建时间"),
        ],
    },
    {
        "title": "（6）联合会诊会话表（consultation_session）",
        "caption": "表 6.6 联合会诊会话表",
        "rows": [
            ("1", "id", "bigint", "20", "√", "主键", "会诊会话主键 ID"),
            ("2", "title", "varchar", "128", "√", "", "会诊标题"),
            ("3", "patient_id", "bigint", "20", "", "外键", "关联患者 ID"),
            ("4", "creator_id", "bigint", "20", "√", "外键", "创建者医生 ID"),
            ("5", "status", "varchar", "32", "√", "", "会诊状态，ACTIVE/CLOSED"),
            ("6", "materials_oss_path", "varchar", "255", "", "", "共享资料包路径"),
            ("7", "created_at", "datetime", "19", "√", "", "创建时间"),
            ("8", "updated_at", "datetime", "19", "√", "", "更新时间"),
        ],
    },
    {
        "title": "（7）审计日志表（audit_log）",
        "caption": "表 6.7 审计日志表",
        "rows": [
            ("1", "id", "bigint", "20", "√", "主键", "日志主键 ID"),
            ("2", "user_id", "bigint", "20", "", "外键", "操作人 ID"),
            ("3", "action", "varchar", "64", "√", "", "操作类型，如 login、review、export"),
            ("4", "detail", "varchar", "255", "", "", "操作详情"),
            ("5", "created_at", "datetime", "19", "√", "", "日志生成时间"),
        ],
    },
]


MOBILE_REQUIREMENT_ROWS = [
    ("终端形态", "支持 Uni-app App-Plus 运行容器的 Android 或 iOS 设备", "Android 医生工作手机、iPad", "项目移动端以 Uni-app 打包发布"),
    ("处理器", "8 核移动处理器", "骁龙 8 系或同等级芯片", "保障列表、上传与 WebView 三维浏览流畅性"),
    ("运行内存", "6GB RAM", "8GB RAM 及以上", "避免大体积影像预览与会诊页面频繁重载"),
    ("存储空间", "预留 2GB 可用空间", "预留 4GB 可用空间", "用于缓存影像切片、附件与调试日志"),
    ("网络环境", "稳定 Wi-Fi 或 5G", "院内千兆 Wi-Fi", "影像上传与模型结果拉取依赖网络质量"),
    ("权限要求", "网络、相机、文件访问权限", "额外启用通知权限", "对应 manifest.json 中的相机与网络权限配置"),
    ("图形能力", "支持 WebGL 的系统 WebView", "支持硬件加速的 WebView", "三维模型页面通过 WebView 加载渲染页面"),
]


DEPLOYMENT_ROWS = [
    ("客户端构建", "Uni-app + Vue 3 + Wot Design Uni", "使用 HBuilderX 或 CLI 打包 App-Plus 产物", "同一套页面支持登录、患者、会诊与 3D 浏览"),
    ("服务地址配置", "前端提供服务地址设置页", "网关地址统一指向 8080 入口", "移动端不直连各微服务，统一经网关路由"),
    ("鉴权接入", "Authorization: Bearer Token", "登录成功后本地缓存 token", "除 /api/auth/* 外，其余接口均需鉴权"),
    ("影像上传", "/api/patients/{id}/studies/{studyId}/upload", "支持 JSON 登记和 multipart 实传", "单次请求大小按后端 500MB 上限控制"),
    ("模型调用", "/api/studies/{studyId}/segment 或 /segment/multimodal", "由 study-service 转发到 model-service", "模型服务默认院内地址为 http://localhost:5001"),
    ("结果展示", "/api/studies/{studyId}/model 与 /artifacts/latest", "轮廓与三维资源写入 OSS 后回传前端", "移动端通过 WebView 展示 NRRD/模型结果"),
]


DETAIL_MODULES = [
    {
        "title": "8.1 账户鉴权与医生管理功能模块",
        "description": (
            "账户鉴权与医生管理模块负责系统用户的注册、登录、令牌签发与身份识别，是整个业务链路的安全入口。"
            "模块以 doctor_user 表为核心，支持医生录入姓名、医院、科室、手机号和密码完成注册；登录成功后由鉴权服务返回 token "
            "以及当前医生基础信息，供前端后续患者管理、影像接入、三维查看和联合会诊页面复用。该模块同时承担账号状态校验职责，"
            "停用账号不能进入业务接口，保证院内部署环境下的账户可控性。"
        ),
        "performance_intro": (
            "本模块的性能目标以“快速登录、轻量校验、低额外开销”为原则。设计状态下，鉴权链路主要发生在网关与 auth-service 之间，"
            "属于短事务请求，响应时间应稳定控制在秒级以内。"
        ),
        "performance_caption": "表 8.1 账户鉴权与医生管理模块系统性能要求表",
        "performance_rows": [
            ("1", "注册账号", "低", "1s 以内"),
            ("2", "账号登录", "高", "500ms 以内"),
            ("3", "令牌校验", "高", "100ms 以内"),
            ("4", "个人信息获取", "中", "300ms 以内"),
        ],
        "input": "注册输入包括姓名、医院、科室、手机号和密码；登录输入包括 username 与 password；个人信息展示输入为当前 token 中的用户标识。",
        "output": (
            "注册输出为新医生主键 ID；登录输出为 token 与医生基础信息；失败时统一输出错误码、失败原因和可读提示信息。"
        ),
        "logic": (
            "医生在移动端提交注册或登录请求后，请求先经过网关路由至 auth-service。注册流程中系统校验手机号唯一性并写入 doctor_user；"
            "登录流程中系统比对密码哈希，校验通过后生成访问令牌并返回给前端。前端将 token 写入本地存储，后续请求通过 Authorization "
            "头透传，网关和业务服务据此完成用户身份恢复。"
        ),
        "limits": "该模块要求手机号唯一，账号状态必须为启用；若 Redis 或网关鉴权链路不可用，则仅允许返回统一失败提示，不允许绕过认证直接访问业务模块。",
    },
    {
        "title": "8.2 患者建档与影像管理功能模块",
        "description": (
            "患者建档与影像管理模块负责患者基础档案、影像序列登记、文件上传和影像元数据维护，是后续 AI 分割与审核流程的数据入口。"
            "模块围绕 patient、patient_study 和 file_upload 三张表组织业务数据，支持患者列表查询、新增患者、编辑患者、登记 study、"
            "上传 DICOM 或 NRRD 文件以及获取文件详情。页面端对应患者列表、患者详情和新增/编辑患者等页面。"
        ),
        "performance_intro": (
            "本模块以稳态录入与大文件可靠传输为目标。患者信息类接口要求快速返回，影像上传类接口则优先保证数据完整性和失败可恢复性。"
        ),
        "performance_caption": "表 8.2 患者建档与影像管理模块系统性能要求表",
        "performance_rows": [
            ("1", "患者列表查询", "高", "500ms 以内"),
            ("2", "新建患者档案", "中", "1s 以内"),
            ("3", "影像序列登记", "中", "1s 以内"),
            ("4", "影像文件登记", "中", "1s 以内"),
            ("5", "大文件上传受理", "中", "2s 内返回受理结果"),
        ],
        "input": (
            "患者建档输入包括姓名、性别、年龄、分期和诊断；影像序列输入包括 studyUid、modality 和描述信息；文件上传输入包括 fileType、"
            "文件大小和文件流或文件路径。"
        ),
        "output": (
            "患者建档输出为患者编号和患者主键；影像登记输出为 studyId 与 studyUid；文件上传输出为 fileId、filePath、fileType 等索引信息。"
        ),
        "logic": (
            "医生先创建患者档案，patient-service 生成 patient_no 并写入 patient 表；随后在患者详情页登记影像序列，study-service 创建 "
            "patient_study 记录。上传阶段可直接提交 multipart 文件，也可先登记后异步上传到对象存储。上传成功后系统写入 file_upload，"
            "并将文件路径与 study 关联，供分割模块继续使用。"
        ),
        "limits": (
            "该模块要求所有写操作携带有效 token；单次上传大小受后端 multipart 配置限制，当前 patient-service 与 study-service 均配置为 500MB。"
            "若对象存储未配置，则仅能保存索引路径或本地占位地址。"
        ),
    },
    {
        "title": "8.3 AI 分割与靶区审核功能模块",
        "description": (
            "AI 分割与靶区审核模块是系统的核心业务模块，负责调用模型服务完成单模态或多模态影像分割，回写轮廓工件，并支撑医生对 GTV、"
            "CTV、OAR 结果进行审核确认。模块涉及 study-service、model-service 以及 contour_result、file_upload 等数据实体。"
            "当前实现支持 /api/studies/{studyId}/segment 和 /segment/multimodal 两条主链路，前者处理单个 NRRD，后者接收 flair、t1、"
            "t1c、t2 四个模态文件。"
        ),
        "performance_intro": (
            "该模块的性能指标分为同步接口指标和离线推理指标两类。同步阶段重点保证提交与状态查询迅速返回；模型推理本身受影像体积和算力"
            "影响，属于分钟级处理任务。"
        ),
        "performance_caption": "表 8.3 AI 分割与靶区审核模块系统性能要求表",
        "performance_rows": [
            ("1", "分割任务提交", "中", "1s 以内返回受理结果"),
            ("2", "轮廓结果查询", "高", "500ms 以内"),
            ("3", "审核状态更新", "高", "300ms 以内"),
            ("4", "单模态分割完成", "中", "分钟级"),
            ("5", "多模态分割完成", "中", "分钟级"),
        ],
        "input": (
            "单模态分割输入为一个合法 NRRD 文件；多模态分割输入为 flair、t1、t1c、t2 四个 NRRD 文件；审核输入为 contourId 与目标状态，"
            "如已确认或待精修。"
        ),
        "output": (
            "分割提交输出为 volumeUrl、labelUrl 以及对应 fileId；轮廓下载输出为 storagePath 与 confidenceMap；审核输出为统一成功或失败提示。"
        ),
        "logic": (
            "医生在移动端选择 study 后触发分割接口，study-service 先校验 study 是否存在，再将体数据上传到对象存储并调用 "
            "model-service。模型返回标签结果后，系统将 volume 和 label 一并落库到 file_upload，同时生成或更新 contour_result 记录。"
            "前端随后调用轮廓查询接口拉取最新结果，医生在审核页面查看、确认或标记待精修，系统再将审核状态与版本信息回写数据库。"
        ),
        "limits": (
            "该模块依赖 model-service.base-url 指向的模型服务可用；多模态接口要求四个模态文件齐全且均为有效 NRRD。若模型推理失败，"
            "系统只能返回失败原因并保留已上传原始体数据，供后续重试。"
        ),
    },
    {
        "title": "8.4 三维可视化与联合会诊功能模块",
        "description": (
            "三维可视化与联合会诊模块支撑医生对分割结果进行空间复核，并围绕病例发起多人协作会诊。三维部分由移动端 model/viewer 页面"
            "通过 WebView 加载 WebGL 页面实现，联合会诊部分由 consultation_session、consultation_member 和 "
            "consultation_message 三张表支撑，支持会诊创建、成员维护、历史消息查询和附件发送。"
        ),
        "performance_intro": (
            "本模块面向高频查看与协作沟通场景，重点保证模型资源索引获取、会诊消息轮询和成员管理操作的实时性。"
        ),
        "performance_caption": "表 8.4 三维可视化与联合会诊模块系统性能要求表",
        "performance_rows": [
            ("1", "三维资源索引获取", "中", "500ms 以内"),
            ("2", "会诊列表查询", "高", "500ms 以内"),
            ("3", "会诊消息轮询", "高", "300ms 以内"),
            ("4", "附件发送", "中", "2s 内返回受理结果"),
            ("5", "成员增删", "低", "500ms 以内"),
        ],
        "input": (
            "三维复核输入为 studyId；会诊创建输入为标题、patientId 与 memberIds；消息输入为文本内容或附件文件，以及 messageType。"
        ),
        "output": (
            "三维接口输出为 modelPath、heatmapPath、oarPath 等资源地址；会诊接口输出为 consultationId、成员列表、消息列表与附件索引路径。"
        ),
        "logic": (
            "前端进入 3D 页面后调用 /api/studies/{studyId}/model 和 /artifacts/latest 获取模型资源，再由 WebView 中的渲染页面解析 NRRD 或"
            "模型文件完成显示。联合会诊流程中，医生在移动端创建 consultation_session，会诊成员写入 consultation_member；"
            "讨论消息和附件通过 socialController 写入 consultation_message，并在前端通过轮询机制持续刷新。"
        ),
        "limits": (
            "该模块要求终端系统 WebView 支持 WebGL；联合会诊的成员添加默认基于好友关系，非好友不能直接加入。若对象存储不可用，则附件"
            "消息只能保留文本说明，无法提供可下载资源。"
        ),
    },
    {
        "title": "8.5 智能助手与知识检索功能模块",
        "description": (
            "智能助手与知识检索功能模块承担项目中的问答辅助与知识支撑能力。前端在会诊中心页面提供 AI 问答入口，调用 "
            "/api/agent/chat 和 /api/agent/chat/stream 与模型侧服务交互，结合本地知识库检索结果为医生提供术语解释、流程指引、"
            "病例资料辅助阅读和会诊记录整理等能力。该模块并不直接修改核心业务数据，而是为医生决策提供参考信息。"
        ),
        "performance_intro": (
            "该模块的性能主要体现为首 token 返回速度、历史消息读取速度和连续对话稳定性。由于其依赖大模型与知识检索链路，系统更关注"
            "流式响应体验而不是一次性全文返回时间。"
        ),
        "performance_caption": "表 8.5 智能助手与知识检索功能模块系统性能要求表",
        "performance_rows": [
            ("1", "会话历史加载", "中", "500ms 以内"),
            ("2", "同步问答接口", "中", "3s 以内返回首轮结果"),
            ("3", "流式问答接口", "高", "2s 内返回首 token"),
            ("4", "知识检索命中", "高", "1s 以内"),
        ],
        "input": (
            "输入包括自然语言问题、可选病例上下文、可选图片或附件，以及当前医生身份标识。"
        ),
        "output": (
            "输出为模型回复内容、流式消息片段、历史对话列表和可能引用的知识片段；失败时返回统一错误信息。"
        ),
        "logic": (
            "医生在会诊中心进入 AI 对话页后，前端先读取历史对话，再根据发送模式调用同步或流式问答接口。模型侧服务根据问题内容执行知识"
            "检索、提示词拼装和大模型推理，将结果持续回传给前端；前端将消息写入会话视图，但不直接写回 patient、study 或 contour 数据表。"
        ),
        "limits": (
            "该模块依赖模型侧服务与知识库可用；回答仅作为辅助建议，不能替代医生最终决策。若知识库检索失败或外部模型不可用，页面应退化为"
            "普通会诊沟通能力。"
        ),
    },
]


def set_east_asia_font(run, font_name: str = FONT_NAME) -> None:
    run.font.name = font_name
    run.font.size = FONT_SIZE
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), font_name)


def strip_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def remove_first_section_break(document: Document) -> None:
    paragraphs = document.paragraphs
    if not paragraphs:
        return
    ppr = paragraphs[0]._p.pPr
    if ppr is not None and ppr.sectPr is not None:
        ppr.remove(ppr.sectPr)


def clone_section_properties(document: Document, template_document: Document) -> None:
    src_sect_pr = template_document.sections[0]._sectPr
    dst_body = document._element.body
    dst_sect_pr = dst_body.sectPr
    if dst_sect_pr is not None:
        dst_body.remove(dst_sect_pr)
    dst_body.append(deepcopy(src_sect_pr))


def add_paragraph(document: Document, text: str = "", style: str = "Normal",
                  alignment: WD_ALIGN_PARAGRAPH | None = None,
                  first_line_indent: Pt | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    if alignment is not None:
        paragraph.alignment = alignment
    if first_line_indent is not None:
        paragraph.paragraph_format.first_line_indent = first_line_indent
    run = paragraph.add_run(text)
    set_east_asia_font(run)


def add_custom_heading(document: Document, text: str, size: Pt) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    run.font.bold = True
    run.font.size = size


def set_cell_text(cell, text: str, align: WD_ALIGN_PARAGRAPH) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    set_east_asia_font(run)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_caption(document: Document, text: str, style: str) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_east_asia_font(run)


def set_table_borders(table) -> None:
    tbl_pr = table._element.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tbl_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def add_db_table(document: Document, caption_style: str, table_info: dict) -> None:
    add_paragraph(document, table_info["title"])
    add_caption(document, table_info["caption"], caption_style)

    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(0.8), Cm(2.8), Cm(1.6), Cm(1.0), Cm(1.2), Cm(1.8), Cm(5.4)]
    headers = ["序号", "字段名", "类型", "长度", "不是 null", "主键/外键", "说明"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = widths[idx]
        set_cell_text(cell, header, WD_ALIGN_PARAGRAPH.CENTER)

    for row_data in table_info["rows"]:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            row.cells[idx].width = widths[idx]
            set_cell_text(row.cells[idx], value, WD_ALIGN_PARAGRAPH.LEFT)

    set_table_borders(table)
    add_paragraph(document, "")


def add_four_col_table(document: Document, caption_style: str, caption: str,
                       headers: list[str], rows: list[tuple[str, str, str, str]],
                       widths: list[Cm]) -> None:
    add_caption(document, caption, caption_style)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = widths[idx]
        set_cell_text(cell, header, WD_ALIGN_PARAGRAPH.CENTER)

    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            row.cells[idx].width = widths[idx]
            set_cell_text(row.cells[idx], value, WD_ALIGN_PARAGRAPH.LEFT)

    set_table_borders(table)
    add_paragraph(document, "")


def build_document() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    template = Document(str(TEMPLATE_PATH))
    body_style_name = template.paragraphs[560].style.name
    caption_style_name = template.paragraphs[521].style.name

    document = Document(str(TEMPLATE_PATH))
    strip_body(document)
    clone_section_properties(document, template)
    remove_first_section_break(document)

    add_custom_heading(document, "6 数据库设计", Pt(16))
    add_paragraph(
        document,
        "结合系统现有 MySQL 脚本、微服务接口和业务流程设计，数据库围绕医生账号、患者档案、影像序列、"
        "AI 分割结果、文件索引、联合会诊和审计日志七类核心实体展开。各实体以 patient_id、study_id、"
        "doctor_id 和 consultation_id 为主线建立关联，支撑“建档-上传-分割-审核-会诊-归档”的闭环业务流程。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_paragraph(
        document,
        "根据当前项目实现，核心数据表详细设计如下；其中 doctor_friend、consultation_member 和 "
        "consultation_message 等扩展协作表用于支撑好友关系和会诊消息，本章以最主要的业务主表为重点展开。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    for table_info in DB_TABLES:
        add_db_table(document, caption_style_name, table_info)

    add_custom_heading(document, "7 手机端侧部署设计", Pt(16))
    add_paragraph(
        document,
        "系统移动端基于 Uni-app 与 Vue 3 开发，承担医生登录、患者管理、影像上传、AI 结果查看、三维浏览和联合会诊等主要业务。"
        "当前项目采用“手机端采集与展示 + 院内后端推理”的协同部署方式，不在移动端本地直接运行分割模型，"
        "从而降低终端算力压力并保证模型版本统一。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_custom_heading(document, "7.1 手机环境需求", Pt(14))
    add_paragraph(
        document,
        "结合 manifest.json 中的运行权限、3D 页面 WebView 渲染方式以及影像上传链路，系统移动端环境建议如下。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_four_col_table(
        document,
        caption_style_name,
        "表 7.1 手机端环境需求表",
        ["项目", "最低要求", "建议配置", "说明"],
        MOBILE_REQUIREMENT_ROWS,
        [Cm(2.2), Cm(4.0), Cm(4.0), Cm(4.4)],
    )

    add_custom_heading(document, "7.2 手机端应用部署", Pt(14))

    add_custom_heading(document, "7.2.1 客户端构建与安装", Pt(12))
    add_paragraph(
        document,
        "移动端客户端由 MasterImage_frontend 工程统一维护，前端页面通过 pages.json 组织登录、患者、会诊、"
        "三维查看等功能入口，并通过 manifest.json 中的 app-plus 配置声明网络、相机等权限。部署时使用 Uni-app 打包 "
        "App-Plus 产物，生成 Android 安装包或 iOS 安装包后分发到院内终端。由于系统提供服务地址设置页，客户端安装后可按医院内部"
        "网络环境配置统一网关地址，无需直接感知各微服务端口。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_custom_heading(document, "7.2.2 影像上传与模型调用部署", Pt(12))
    add_paragraph(
        document,
        "手机端完成登录后，所有业务请求均经网关 8080 统一转发。影像登记与上传通过 /api/patients/{id}/studies/* "
        "接口进入 study-service，文件写入 OSS 后形成 file_upload 记录；当医生在移动端触发 AI 分割时，study-service 再通过 "
        "ModelSegmentService 调用院内 model-service，默认服务地址为 http://localhost:5001。对于多模态分割，"
        "前端需一次性提交 flair、t1、t1c、t2 四类 NRRD 文件，后端仅在四个文件均合法时发起推理请求。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_custom_heading(document, "7.2.3 结果回传与三维展示部署", Pt(12))
    add_paragraph(
        document,
        "分割结果生成后，标签文件与体数据统一回传到对象存储，并通过 /api/studies/{studyId}/model、"
        "/api/studies/{studyId}/artifacts/latest 以及下载接口提供给前端。移动端的 3D 查看页面通过 WebView 加载渲染内容，"
        "在终端侧主要负责资源下载、阈值与透明度控制、交互展示，不承担重型模型推理任务。这样既保证了移动端操作体验，"
        "也使模型升级仅需在院内服务侧完成一次发布即可全端生效。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_four_col_table(
        document,
        caption_style_name,
        "表 7.2 手机端部署链路表",
        ["环节", "当前实现", "部署要求", "备注"],
        DEPLOYMENT_ROWS,
        [Cm(2.2), Cm(4.0), Cm(4.0), Cm(4.4)],
    )

    add_custom_heading(document, "8 详细设计", Pt(16))
    add_paragraph(
        document,
        "结合当前项目已有前端页面、微服务控制器、数据库脚本和模型侧服务，系统详细设计按账户鉴权、患者与影像、AI 分割审核、"
        "三维会诊和智能助手五个核心功能模块展开。",
        style=body_style_name,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    perf_widths = [Cm(1.2), Cm(6.0), Cm(3.0), Cm(4.0)]
    for module in DETAIL_MODULES:
        add_custom_heading(document, module["title"], Pt(14))

        add_custom_heading(document, module["title"].split(" ")[0] + ".1 功能描述", Pt(12))
        add_paragraph(document, module["description"], style=body_style_name, alignment=WD_ALIGN_PARAGRAPH.LEFT)

        add_custom_heading(document, module["title"].split(" ")[0] + ".2 性能描述", Pt(12))
        add_paragraph(document, module["performance_intro"], style=body_style_name, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        add_four_col_table(
            document,
            caption_style_name,
            module["performance_caption"],
            ["序号", "功能", "用户使用频率", "要求响应时间"],
            module["performance_rows"],
            perf_widths,
        )

        add_custom_heading(document, module["title"].split(" ")[0] + ".3 输入", Pt(12))
        add_paragraph(document, module["input"], style=body_style_name, alignment=WD_ALIGN_PARAGRAPH.LEFT)

        add_custom_heading(document, module["title"].split(" ")[0] + ".4 输出", Pt(12))
        add_paragraph(document, module["output"], style=body_style_name, alignment=WD_ALIGN_PARAGRAPH.LEFT)

        add_custom_heading(document, module["title"].split(" ")[0] + ".5 程序逻辑", Pt(12))
        add_paragraph(document, module["logic"], style=body_style_name, alignment=WD_ALIGN_PARAGRAPH.LEFT)

        add_custom_heading(document, module["title"].split(" ")[0] + ".6 限制条件", Pt(12))
        add_paragraph(document, module["limits"], style=body_style_name, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    document.save(str(OUTPUT_PATH))


if __name__ == "__main__":
    build_document()
