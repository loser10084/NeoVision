from __future__ import annotations

import json
import shutil
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


ROOT = Path(r"D:\project\AI\ImageIdentification")
DOC_DIR = ROOT / "output" / "doc"
SOURCE = DOC_DIR / "灵犀智影-项目开发文档_6_7_8_排版修复版.docx"
OUTPUT = DOC_DIR / "灵犀智影-项目开发文档_6_7_8_排版修复版_TimesNewRoman.docx"
REPORT = DOC_DIR / "灵犀智影-项目开发文档_6_7_8_排版修复版_TimesNewRoman_报告.json"
TMP = DOC_DIR / "._tmp_6_7_8_tnr.docx"

FONT = "Times New Roman"


def set_run_fonts(run) -> None:
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:cs"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)


def normalize_style(style, size: Pt | None = None, bold: bool | None = None) -> None:
    style.font.name = FONT
    if size is not None:
        style.font.size = size
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:cs"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)


def clear_run_conflicts(run) -> int:
    removed = 0
    rpr = run._r.rPr
    if rpr is None:
        return 0
    for tag in ("w:rFonts", "w:rStyle", "w:lang", "w:noProof"):
        elem = rpr.find(qn(tag))
        if elem is not None:
            rpr.remove(elem)
            removed += 1
    return removed


def normalize_paragraph_runs(paragraph, report: dict) -> None:
    if not paragraph.text.strip():
        return
    for run in paragraph.runs:
        if not run.text:
            continue
        report["removed_run_conflicts"] += clear_run_conflicts(run)
        size = run.font.size
        bold = run.bold
        italic = run.italic
        underline = run.underline
        set_run_fonts(run)
        run.font.size = size
        run.bold = bold
        run.italic = italic
        run.underline = underline
        report["normalized_runs"] += 1


def normalize_doc_defaults(docx_path: Path) -> None:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(docx_path, "r") as zin, ZipFile(TMP, "w", compression=ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/styles.xml":
                root = etree.fromstring(data)
                defaults = root.xpath(".//w:docDefaults/w:rPrDefault/w:rPr", namespaces=ns)
                if defaults:
                    rpr = defaults[0]
                    rfonts = rpr.find(qn("w:rFonts"))
                    if rfonts is None:
                        rfonts = OxmlElement("w:rFonts")
                        rpr.insert(0, rfonts)
                    rfonts.attrib.clear()
                    rfonts.set(qn("w:ascii"), FONT)
                    rfonts.set(qn("w:hAnsi"), FONT)
                    rfonts.set(qn("w:cs"), FONT)
                    rfonts.set(qn("w:eastAsia"), FONT)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            zout.writestr(item, data)
    shutil.move(str(TMP), str(docx_path))


def main() -> None:
    shutil.copyfile(SOURCE, OUTPUT)
    doc = Document(str(OUTPUT))

    report = {
        "source": str(SOURCE),
        "output": str(OUTPUT),
        "font_target": FONT,
        "removed_run_conflicts": 0,
        "normalized_runs": 0,
        "normalized_paragraphs": 0,
        "normalized_table_paragraphs": 0,
    }

    normalize_style(doc.styles["Normal"], Pt(10.5), None)
    normalize_style(doc.styles["Heading 1"], Pt(16), True)
    normalize_style(doc.styles["Heading 2"], Pt(14), True)
    normalize_style(doc.styles["Heading 3"], Pt(12), True)

    for paragraph in doc.paragraphs:
        if not paragraph.text.strip():
            continue
        normalize_paragraph_runs(paragraph, report)
        report["normalized_paragraphs"] += 1

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if not paragraph.text.strip():
                        continue
                    normalize_paragraph_runs(paragraph, report)
                    report["normalized_table_paragraphs"] += 1

    doc.save(str(OUTPUT))
    normalize_doc_defaults(OUTPUT)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
