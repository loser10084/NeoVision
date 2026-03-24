from __future__ import annotations

import json
import re
import shutil
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


ROOT = Path(r"D:\project\AI\ImageIdentification")
SOURCE_PATH = ROOT / "test" / "灵犀智影-项目开发文档.docx"
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_PATH = OUTPUT_DIR / "灵犀智影-项目开发文档_排版修复版.docx"
REPORT_PATH = OUTPUT_DIR / "灵犀智影-项目开发文档_排版修复报告.json"
TMP_XML_PATH = OUTPUT_DIR / "._tmp_layout_cleanup.docx"

WEST_FONT = "Times New Roman"
BODY_EAST_FONT = "宋体"
HEADING_EAST_FONT = "黑体"
BODY_SIZE = Pt(12)
CAPTION_SIZE = Pt(11)
TABLE_SIZE = Pt(10.5)

SPECIAL_CHAR_MAP = {
    "\u00a0": " ",
    "\u2000": " ",
    "\u2001": " ",
    "\u2002": " ",
    "\u2003": " ",
    "\u2004": " ",
    "\u2005": " ",
    "\u2006": " ",
    "\u2007": " ",
    "\u2008": " ",
    "\u2009": " ",
    "\u200a": " ",
    "\u202f": " ",
    "\u200b": "",
    "\u200c": "",
    "\u200d": "",
    "\ufeff": "",
    "\v": " ",
}

TERMINAL_PUNCT = ("。", "！", "？", "；", "：", ".", "!", "?", ";", ":", "）", "】", "」", "”")
NON_MERGE_PREFIX = re.compile(r"^(?:[①②③④⑤⑥⑦⑧⑨⑩]|（\d+）|\(\d+\)|图\s*\d|表\s*\d|组件注解：|4(?:\.\d+)+)")


def clean_text(text: str) -> tuple[str, int]:
    count = 0
    result = text
    for src, dst in SPECIAL_CHAR_MAP.items():
        if src in result:
            count += result.count(src)
            result = result.replace(src, dst)
    result = re.sub(r"[ ]{2,}", " ", result)
    return result, count


def set_rfonts(font_obj, east_font: str, west_font: str = WEST_FONT) -> None:
    font_obj.name = west_font
    rpr = font_obj._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = etree.Element(qn("w:rFonts"))
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), west_font)
    rfonts.set(qn("w:hAnsi"), west_font)
    rfonts.set(qn("w:cs"), west_font)
    rfonts.set(qn("w:eastAsia"), east_font)


def normalize_style(style, east_font: str, size: Pt | None = None, bold: bool | None = None) -> None:
    set_rfonts(style.font, east_font)
    if size is not None:
        style.font.size = size
    if bold is not None:
        style.font.bold = bold


def clear_run_conflicts(run) -> int:
    removed = 0
    rpr = run._r.rPr
    if rpr is None:
        return 0
    for tag in ("w:rFonts", "w:rStyle", "w:lang", "w:noProof"):
        elem = rpr.find(qn(tag))
        if elem is not None:
            rpr.remove(elem)
            removed += 1
    return removed


def is_caption(paragraph) -> bool:
    txt = paragraph.text.strip()
    return paragraph.alignment == 1 and (txt.startswith("图 ") or txt.startswith("表 "))


def normalize_paragraph_runs(paragraph, east_font: str, default_size: Pt, report: dict) -> None:
    for run in paragraph.runs:
        report["removed_run_conflicts"] += clear_run_conflicts(run)
        if run.text:
            new_text, count = clean_text(run.text)
            if count:
                run.text = new_text
                report["special_char_replacements"] += count
        set_rfonts(run.font, east_font)
        if run.font.size is None:
            run.font.size = default_size
        report["normalized_runs"] += 1


def paragraph_text(paragraph) -> str:
    return paragraph.text.strip()


def can_merge(a, b) -> bool:
    ta = paragraph_text(a)
    tb = paragraph_text(b)
    if not ta or not tb:
        return False
    if a.style.name != "Normal" or b.style.name != "Normal":
        return False
    if a.alignment == 1 or b.alignment == 1:
        return False
    if NON_MERGE_PREFIX.match(tb):
        return False
    if re.match(r"^(?:[①②③④⑤⑥⑦⑧⑨⑩]|（\d+）|\(\d+\))", ta):
        return False
    if ta.startswith("图 ") or ta.startswith("表 ") or tb.startswith("图 ") or tb.startswith("表 "):
        return False
    if ta.endswith(TERMINAL_PUNCT):
        return False
    if re.match(r"^[A-Z0-9①②③④⑤⑥⑦⑧⑨⑩（(]", tb):
        return False
    if len(ta) < 6 or len(tb) < 3:
        return False
    return True


def append_paragraph_text(target, source) -> None:
    joiner = ""
    if target.text and source.text:
        last_char = target.text[-1]
        first_char = source.text[0]
        if last_char.isascii() and first_char.isascii():
            joiner = " "
    target.add_run(joiner + source.text)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def merge_paragraphs(document: Document, start_index: int, report: dict) -> None:
    idx = start_index
    while idx < len(document.paragraphs) - 1:
        current = document.paragraphs[idx]
        nxt = document.paragraphs[idx + 1]
        if can_merge(current, nxt):
            report["merged_paragraph_pairs"] += 1
            append_paragraph_text(current, nxt)
            delete_paragraph(nxt)
            continue
        idx += 1


def remove_xml_noise(docx_path: Path, report: dict) -> None:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(docx_path, "r") as zin, ZipFile(TMP_XML_PATH, "w", compression=ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                proof_err = root.xpath(".//w:proofErr", namespaces=ns)
                no_proof = root.xpath(".//w:noProof", namespaces=ns)
                for elem in proof_err:
                    elem.getparent().remove(elem)
                for elem in no_proof:
                    elem.getparent().remove(elem)
                report["removed_proof_err"] = len(proof_err)
                report["removed_no_proof"] = len(no_proof)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            zout.writestr(item, data)
    shutil.move(str(TMP_XML_PATH), str(docx_path))


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(SOURCE_PATH, OUTPUT_PATH)
    document = Document(str(OUTPUT_PATH))

    report = {
        "source": str(SOURCE_PATH),
        "output": str(OUTPUT_PATH),
        "special_char_replacements": 0,
        "removed_run_conflicts": 0,
        "normalized_runs": 0,
        "normalized_body_paragraphs": 0,
        "normalized_heading_paragraphs": 0,
        "normalized_table_paragraphs": 0,
        "merged_paragraph_pairs": 0,
        "removed_proof_err": 0,
        "removed_no_proof": 0,
        "unsafe_merge_positions": [],
    }

    normalize_style(document.styles["Normal"], BODY_EAST_FONT, BODY_SIZE, None)
    normalize_style(document.styles["Heading 1"], HEADING_EAST_FONT, Pt(16), True)
    normalize_style(document.styles["Heading 2"], HEADING_EAST_FONT, Pt(14), True)
    normalize_style(document.styles["Heading 3"], HEADING_EAST_FONT, Pt(12), True)

    main_start = 0
    for i, paragraph in enumerate(document.paragraphs):
        if paragraph.style.name == "Heading 1":
            main_start = i
            break

    merge_paragraphs(document, main_start, report)

    for idx, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph)
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            normalize_paragraph_runs(paragraph, HEADING_EAST_FONT, {
                "Heading 1": Pt(16),
                "Heading 2": Pt(14),
                "Heading 3": Pt(12),
            }.get(paragraph.style.name, Pt(12)), report)
            report["normalized_heading_paragraphs"] += 1
        else:
            size = CAPTION_SIZE if is_caption(paragraph) else BODY_SIZE
            normalize_paragraph_runs(paragraph, BODY_EAST_FONT, size, report)
            report["normalized_body_paragraphs"] += 1

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if not paragraph_text(paragraph):
                        continue
                    normalize_paragraph_runs(paragraph, BODY_EAST_FONT, TABLE_SIZE, report)
                    report["normalized_table_paragraphs"] += 1

    document.save(str(OUTPUT_PATH))
    remove_xml_noise(OUTPUT_PATH, report)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
