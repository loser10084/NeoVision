from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


ROOT = Path(r"D:\project\AI\ImageIdentification")
SOURCE = ROOT / "output" / "doc" / "灵犀智影-项目开发文档_6_7_8.docx"
OUT_DIR = ROOT / "output" / "doc"
OUTPUT = OUT_DIR / "灵犀智影-项目开发文档_6_7_8_排版修复版.docx"
REPORT = OUT_DIR / "灵犀智影-项目开发文档_6_7_8_排版修复报告.json"
TMP = OUT_DIR / "._tmp_6_7_8_cleanup.docx"

WEST_FONT = "Times New Roman"
BODY_EAST_FONT = "宋体"
HEADING_EAST_FONT = "黑体"
BODY_SIZE = Pt(10.5)
CAPTION_SIZE = Pt(11)
TABLE_SIZE = Pt(10.5)

SPECIAL_CHAR_MAP = {
    "\u00a0": " ",
    "\u2000": " ",
    "\u2001": " ",
    "\u2002": " ",
    "\u2003": " ",
    "\u2004": " ",
    "\u2005": " ",
    "\u2006": " ",
    "\u2007": " ",
    "\u2008": " ",
    "\u2009": " ",
    "\u200a": " ",
    "\u202f": " ",
    "\u200b": "",
    "\u200c": "",
    "\u200d": "",
    "\ufeff": "",
    "\v": " ",
}


def clean_text(text: str) -> tuple[str, int]:
    count = 0
    out = text
    for src, dst in SPECIAL_CHAR_MAP.items():
        if src in out:
            count += out.count(src)
            out = out.replace(src, dst)
    out = re.sub(r"[ ]{2,}", " ", out)
    return out, count


def set_run_fonts(run, east_font: str, west_font: str = WEST_FONT) -> None:
    run.font.name = west_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), west_font)
    rfonts.set(qn("w:hAnsi"), west_font)
    rfonts.set(qn("w:cs"), west_font)
    rfonts.set(qn("w:eastAsia"), east_font)


def normalize_style_fonts(style, east_font: str, size: Pt | None = None, bold: bool | None = None) -> None:
    style.font.name = WEST_FONT
    style.font.size = size or style.font.size
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), WEST_FONT)
    rfonts.set(qn("w:hAnsi"), WEST_FONT)
    rfonts.set(qn("w:cs"), WEST_FONT)
    rfonts.set(qn("w:eastAsia"), east_font)


def clear_run_conflicts(run) -> int:
    removed = 0
    rpr = run._r.rPr
    if rpr is None:
        return 0
    for tag in ("w:rFonts", "w:rStyle", "w:lang", "w:noProof"):
        elem = rpr.find(qn(tag))
        if elem is not None:
            rpr.remove(elem)
            removed += 1
    return removed


def is_caption(paragraph) -> bool:
    txt = paragraph.text.strip()
    return paragraph.alignment == 1 and (txt.startswith("表 ") or txt.startswith("图 "))


def is_heading_like(paragraph) -> bool:
    txt = paragraph.text.strip()
    if paragraph.style.name.startswith("Heading"):
        return True
    return bool(re.match(r"^(?:\d+(?:\.\d+){0,2}\s+)", txt))


def normalize_paragraph(paragraph, report: dict) -> None:
    txt = paragraph.text.strip()
    if not txt:
        return
    east_font = HEADING_EAST_FONT if is_heading_like(paragraph) else BODY_EAST_FONT
    default_size = CAPTION_SIZE if is_caption(paragraph) else BODY_SIZE
    if is_heading_like(paragraph):
        report["normalized_heading_paragraphs"] += 1
    else:
        report["normalized_body_paragraphs"] += 1
    for run in paragraph.runs:
        if not run.text:
            continue
        report["removed_run_conflicts"] += clear_run_conflicts(run)
        cleaned, repl = clean_text(run.text)
        if repl:
            run.text = cleaned
            report["special_char_replacements"] += repl
        original_size = run.font.size
        original_bold = run.bold
        original_italic = run.italic
        original_underline = run.underline
        set_run_fonts(run, east_font)
        run.font.size = original_size or default_size
        run.bold = original_bold
        run.italic = original_italic
        run.underline = original_underline
        report["normalized_runs"] += 1


def normalize_table_fonts(document: Document, report: dict) -> None:
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if not paragraph.text.strip():
                        continue
                    report["normalized_table_paragraphs"] += 1
                    for run in paragraph.runs:
                        if not run.text:
                            continue
                        report["removed_run_conflicts"] += clear_run_conflicts(run)
                        cleaned, repl = clean_text(run.text)
                        if repl:
                            run.text = cleaned
                            report["special_char_replacements"] += repl
                        original_size = run.font.size
                        original_bold = run.bold
                        original_italic = run.italic
                        original_underline = run.underline
                        set_run_fonts(run, BODY_EAST_FONT)
                        run.font.size = original_size or TABLE_SIZE
                        run.bold = original_bold
                        run.italic = original_italic
                        run.underline = original_underline
                        report["normalized_runs"] += 1


def remove_xml_noise(docx_path: Path, report: dict) -> None:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(docx_path, "r") as zin, ZipFile(TMP, "w", compression=ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                proof_err = root.xpath(".//w:proofErr", namespaces=ns)
                no_proof = root.xpath(".//w:noProof", namespaces=ns)
                for elem in proof_err:
                    elem.getparent().remove(elem)
                for elem in no_proof:
                    elem.getparent().remove(elem)
                report["removed_proof_err"] = len(proof_err)
                report["removed_no_proof"] = len(no_proof)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            elif item.filename == "word/styles.xml":
                root = etree.fromstring(data)
                nsw = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                defaults = root.xpath(".//w:docDefaults/w:rPrDefault/w:rPr", namespaces=nsw)
                if defaults:
                    rpr = defaults[0]
                    rfonts = rpr.find(qn("w:rFonts"))
                    if rfonts is None:
                        rfonts = OxmlElement("w:rFonts")
                        rpr.insert(0, rfonts)
                    rfonts.attrib.clear()
                    rfonts.set(qn("w:ascii"), WEST_FONT)
                    rfonts.set(qn("w:hAnsi"), WEST_FONT)
                    rfonts.set(qn("w:cs"), WEST_FONT)
                    rfonts.set(qn("w:eastAsia"), BODY_EAST_FONT)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            zout.writestr(item, data)
    shutil.move(str(TMP), str(docx_path))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(SOURCE, OUTPUT)
    doc = Document(str(OUTPUT))

    report = {
        "source": str(SOURCE),
        "output": str(OUTPUT),
        "special_char_replacements": 0,
        "removed_run_conflicts": 0,
        "normalized_runs": 0,
        "normalized_body_paragraphs": 0,
        "normalized_heading_paragraphs": 0,
        "normalized_table_paragraphs": 0,
        "removed_proof_err": 0,
        "removed_no_proof": 0,
        "merged_paragraph_pairs": 0,
        "unsafe_merge_positions": [],
    }

    normalize_style_fonts(doc.styles["Normal"], BODY_EAST_FONT, BODY_SIZE, None)
    normalize_style_fonts(doc.styles["Heading 1"], HEADING_EAST_FONT, Pt(16), True)
    normalize_style_fonts(doc.styles["Heading 2"], HEADING_EAST_FONT, Pt(14), True)
    normalize_style_fonts(doc.styles["Heading 3"], HEADING_EAST_FONT, Pt(12), True)

    for paragraph in doc.paragraphs:
        normalize_paragraph(paragraph, report)
    normalize_table_fonts(doc, report)

    doc.save(str(OUTPUT))
    remove_xml_noise(OUTPUT, report)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
